import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from typing import List, Dict, Any, Optional

def parse_all_markdown_tables(text: str) -> List[Dict[str, Any]]:
    lines = text.split("\n")
    tables = []
    current_table = None

    for line in lines:
        stripped = line.strip()
        if stripped.startswith("|"):
            parts = [p.strip() for p in stripped.split("|")[1:-1]]
            # Bỏ qua dòng phân cách (ví dụ |:---|:---|)
            if parts and all(re.match(r"^:?-+:?$", p) for p in parts):
                continue
            if not any(parts):
                continue

            if current_table is None:
                current_table = {"headers": parts, "rows": []}
            else:
                current_table["rows"].append(parts)
        else:
            if current_table:
                tables.append(current_table)
                current_table = None

    if current_table:
        tables.append(current_table)

    return tables

def create_excel_from_tables(tables: List[Dict[str, Any]], filepath: str, query: str = ""):
    wb = openpyxl.Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)

    has_sheets = False

    # Styling UAT
    header_font = Font(name="Times New Roman", size=14, bold=True, color="000000")
    header_fill = PatternFill(start_color="93C47D", end_color="93C47D", fill_type="solid")
    group_font = Font(name="Times New Roman", size=14, bold=True, color="000000")
    group_fill = PatternFill(start_color="C9DAF8", end_color="C9DAF8", fill_type="solid")
    uat_col_a_empty_font = Font(name="Arial", size=10, bold=False, color="000000")
    uat_data_font = Font(name="Times New Roman", size=14, bold=False, color="000000")

    thin_border = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000"),
    )

    # Styling Use Case
    uc_title_font = Font(name="Arial", size=14, bold=True, color="1F4E78")
    uc_subtitle_font = Font(name="Arial", size=10, bold=False, color="595959")
    uc_header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    uc_header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    uc_data_font = Font(name="Arial", size=10, bold=False, color="000000")
    uc_data_font_bold = Font(name="Arial", size=10, bold=True, color="000000")
    uc_group_fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")

    uc_thin_border = Border(
        left=Side(style="thin", color="BFBFBF"),
        right=Side(style="thin", color="BFBFBF"),
        top=Side(style="thin", color="BFBFBF"),
        bottom=Side(style="thin", color="BFBFBF"),
    )

    data_font = Font(name="Arial", size=11, bold=False, color="000000")

    for idx, table in enumerate(tables):
        headers = table["headers"]
        rows = table["rows"]
        if not rows:
            continue

        if len(headers) == 5 and any("Sub Module" in str(h) or "Mô tả" in str(h) for h in headers):
            # UAT Testcase (5 cột MD -> 9 cột Excel)
            ws = wb.create_sheet(title="Kịch bản UAT")
            ws.views.sheetView[0].showGridLines = True

            custom_headers = [
                "Sub Module", "ID", "Mô tả", "Bước thực hiện",
                "Dữ liệu test", "Kết quả mong đợi", "Trạng thái (web)",
                "Kết quả thực tế", "Ghi chú"
            ]
            for c_idx, h_text in enumerate(custom_headers, 1):
                cell = ws.cell(row=1, column=c_idx, value=h_text)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = thin_border

            prev_group = None
            for r_idx, row in enumerate(rows, 2):
                while len(row) < 5:
                    row.append("")

                md_sub_module, md_mota, md_steps, md_expected, md_test_data = row[:5]
                current_group = str(md_sub_module).strip() if md_sub_module else ""
                is_group_change = (current_group != prev_group)
                if is_group_change:
                    prev_group = current_group
                    cell_sub_module_val = md_sub_module
                else:
                    cell_sub_module_val = ""

                excel_row_vals = [
                    cell_sub_module_val,
                    f'=IF(G{r_idx}="","",COUNTA($G$2:G{r_idx}))',
                    md_mota,
                    md_steps,
                    md_test_data,
                    md_expected,
                    "Pass",
                    "",
                    "",
                ]

                for c_idx, val in enumerate(excel_row_vals, 1):
                    cleaned_val = str(val).replace("<br>", "\n").replace("<br/>", "\n").replace("<br />", "\n")
                    cell = ws.cell(row=r_idx, column=c_idx, value=cleaned_val)
                    if c_idx == 1:
                        if is_group_change:
                            cell.font = group_font
                            cell.fill = group_fill
                        else:
                            cell.font = uat_col_a_empty_font
                    else:
                        cell.font = uat_data_font

                    if c_idx in [1, 2, 7, 8]:
                        cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
                    else:
                        cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                    cell.border = thin_border

            col_widths = {
                "A": 25, "B": 8, "C": 35, "D": 40,
                "E": 25, "F": 40, "G": 15, "H": 15, "I": 20
            }
            for col_letter, width in col_widths.items():
                ws.column_dimensions[col_letter].width = width
            has_sheets = True

        elif len(headers) == 6:
            # Bảng Tổng quan UC
            ws = wb.create_sheet(title="Tổng quan UC")
            ws.views.sheetView[0].showGridLines = True

            title_text = "DANH SÁCH USE CASE"
            if query:
                q_clean = re.sub(r"(?i)^/?(usecase|testcase)\s*", "", query).strip()
                if q_clean:
                    title_text += f" — {q_clean.upper()}"
            ws.cell(row=1, column=1, value=title_text).font = uc_title_font
            ws.cell(row=2, column=1, value="Sinh tự động từ tài liệu SRS").font = uc_subtitle_font

            for c_idx, h_text in enumerate(headers, 1):
                cell = ws.cell(row=4, column=c_idx, value=h_text)
                cell.font = uc_header_font
                cell.fill = uc_header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = uc_thin_border

            prev_group = None
            for r_idx, row in enumerate(rows, 5):
                while len(row) < 6:
                    row.append("")

                current_group = str(row[1]).strip() if row[1] else ""
                is_group_change = (current_group != prev_group)
                if is_group_change:
                    prev_group = current_group

                for c_idx, val in enumerate(row, 1):
                    cleaned_val = str(val).replace("<br>", "\n").replace("<br/>", "\n").replace("<br />", "\n")
                    cell = ws.cell(row=r_idx, column=c_idx, value=cleaned_val)
                    cell.font = uc_data_font_bold if is_group_change else uc_data_font
                    if is_group_change:
                        cell.fill = uc_group_fill
                    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                    cell.border = uc_thin_border

            col_widths = {"A": 12, "B": 26, "C": 42, "D": 24, "E": 25, "F": 30}
            for col_letter, width in col_widths.items():
                ws.column_dimensions[col_letter].width = width
            has_sheets = True

        elif len(headers) == 10:
            # Bảng Chi tiết UC
            ws = wb.create_sheet(title="Chi tiết UC")
            ws.views.sheetView[0].showGridLines = True
            ws.cell(row=1, column=1, value="CHI TIẾT TỪNG USE CASE").font = uc_title_font

            for c_idx, h_text in enumerate(headers, 1):
                cell = ws.cell(row=3, column=c_idx, value=h_text)
                cell.font = uc_header_font
                cell.fill = uc_header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = uc_thin_border

            prev_group = None
            for r_idx, row in enumerate(rows, 4):
                while len(row) < 10:
                    row.append("")

                current_group = str(row[1]).strip() if row[1] else ""
                is_group_change = (current_group != prev_group)
                if is_group_change:
                    prev_group = current_group

                for c_idx, val in enumerate(row, 1):
                    cleaned_val = str(val).replace("<br>", "\n").replace("<br/>", "\n").replace("<br />", "\n")
                    cell = ws.cell(row=r_idx, column=c_idx, value=cleaned_val)
                    cell.font = uc_data_font
                    if is_group_change:
                        cell.fill = uc_group_fill
                    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                    cell.border = uc_thin_border

            col_widths = {
                "A": 12, "B": 22, "C": 32, "D": 22, "E": 26,
                "F": 22, "G": 60, "H": 40, "I": 28, "J": 15
            }
            for col_letter, width in col_widths.items():
                ws.column_dimensions[col_letter].width = width
            has_sheets = True

        else:
            # Generic table
            title = f"Sheet {idx+1}"
            ws = wb.create_sheet(title=title)
            ws.views.sheetView[0].showGridLines = True

            for c_idx, h_text in enumerate(headers, 1):
                cell = ws.cell(row=1, column=c_idx, value=h_text)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = thin_border

            for r_idx, row in enumerate(rows, 2):
                while len(row) < len(headers):
                    row.append("")
                for c_idx, val in enumerate(row, 1):
                    cleaned_val = str(val).replace("<br>", "\n").replace("<br/>", "\n").replace("<br />", "\n")
                    cell = ws.cell(row=r_idx, column=c_idx, value=cleaned_val)
                    cell.font = data_font
                    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                    cell.border = thin_border

            for col in ws.columns:
                max_len = 0
                col_letter = col[0].column_letter
                for cell in col:
                    if cell.value:
                        for line in str(cell.value).split("\n"):
                            if len(line) > max_len:
                                max_len = len(line)
                ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)
            has_sheets = True

    if not has_sheets:
        ws = wb.create_sheet(title="Dữ liệu")
        ws.views.sheetView[0].showGridLines = True

    wb.save(filepath)
